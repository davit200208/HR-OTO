import base64
import csv
import io
import uuid
from datetime import UTC, date, datetime

from fastapi import APIRouter, Body, Depends, HTTPException, Query, Request
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.api.dependencies import AuthContext, require_tenant
from app.core.tenant_security import enforce_tenant_access
from app.db.session import get_db
from app.models.admin import Branch, Department, Employee, EmployeeDocument, Position, TenantSetting
from app.models.identity import Tenant, TenantStatus, User, UserRole
from app.services.audit import write_audit
from app.services.entitlements import effective_entitlements, require_feature, require_limit
from app.services.tenant_provisioning import tenant_url

router = APIRouter(prefix="/admin", tags=["company-admin"])

class NamedCreate(BaseModel): name: str = Field(min_length=2, max_length=120)
class PositionCreate(NamedCreate): department_id: uuid.UUID | None = None
class BranchCreate(BaseModel): name: str = Field(min_length=2, max_length=120); address: str = ""; manager_user_id: uuid.UUID | None = None
class EmployeeInput(BaseModel):
    employee_id: str = Field(min_length=1, max_length=48); first_name: str = Field(min_length=1, max_length=80); last_name: str = Field(min_length=1, max_length=80)
    email: str = ""; phone: str = ""; personal_id: str = ""; birth_date: date | None = None; department_id: uuid.UUID | None = None; position_id: uuid.UUID | None = None; branch_id: uuid.UUID | None = None
    employment_date: date | None = None; employment_status: str = "active"; salary_amount: int | None = None; bonus_amount: int | None = None; salary_currency: str = "GEL"
class SettingInput(BaseModel): key: str = Field(min_length=2, max_length=100); value: str = ""
class RoleUpdate(BaseModel): role: UserRole
class PhotoInput(BaseModel): content_base64: str = Field(min_length=4); content_type: str = "image/jpeg"

def _admin(context: AuthContext) -> None:
    if context.user.role not in {UserRole.COMPANY_OWNER, UserRole.HR_ADMIN, UserRole.MANAGER}:
        raise HTTPException(403, "Company Admin permission required")

async def _ctx(company_code: str, context: AuthContext, db: AsyncSession):
    tenant = await db.scalar(select(Tenant).where(Tenant.company_code == company_code.upper(), Tenant.status == TenantStatus.ACTIVE))
    if not tenant: raise HTTPException(404, "Company not found")
    enforce_tenant_access(context.user.tenant_id, tenant.id)
    engine = create_async_engine(tenant_url(tenant.database_name)); return tenant, engine, async_sessionmaker(engine, expire_on_commit=False)

def _employee(e: Employee, include_sensitive: bool = False):
    data = {"id": str(e.id), "employee_id": e.employee_id, "first_name": e.first_name, "last_name": e.last_name, "email": e.email, "phone": e.phone, "department_id": str(e.department_id) if e.department_id else None, "position_id": str(e.position_id) if e.position_id else None, "branch_id": str(e.branch_id) if e.branch_id else None, "employment_date": e.employment_date, "employment_status": e.employment_status, "archived": bool(e.archived_at)}
    if include_sensitive: data.update({"personal_id": e.personal_id, "birth_date": e.birth_date, "salary_amount": e.salary_amount, "bonus_amount": e.bonus_amount, "salary_currency": e.salary_currency, "has_photo": e.photo_data is not None})
    return data

@router.get("/{company_code}/dashboard")
async def dashboard(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); tenant, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            total = await tdb.scalar(select(func.count(Employee.id)).where(Employee.archived_at == "")) or 0
            archived = await tdb.scalar(select(func.count(Employee.id)).where(Employee.archived_at != "")) or 0
            branches = await tdb.scalar(select(func.count(Branch.id)).where(Branch.is_active.is_(True))) or 0
            departments = await tdb.scalar(select(func.count(Department.id)).where(Department.is_active.is_(True))) or 0
            return {"company": tenant.name, "employees": total, "archived_employees": archived, "branches": branches, "departments": departments}
    finally: await engine.dispose()

@router.get("/{company_code}/employees")
async def employees(company_code: str, search: str = "", status: str = "active", context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            query = select(Employee).order_by(Employee.last_name, Employee.first_name)
            if status == "active": query = query.where(Employee.archived_at == "")
            elif status == "archived": query = query.where(Employee.archived_at != "")
            if search: query = query.where(or_(Employee.first_name.ilike(f"%{search}%"), Employee.last_name.ilike(f"%{search}%"), Employee.employee_id.ilike(f"%{search}%"), Employee.email.ilike(f"%{search}%")))
            return [_employee(x) for x in (await tdb.scalars(query)).all()]
    finally: await engine.dispose()

@router.post("/{company_code}/employees", status_code=201)
async def create_employee(company_code: str, payload: EmployeeInput, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); tenant, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            current = await tdb.scalar(select(func.count(Employee.id)).where(Employee.archived_at == "")); await require_limit(db, tenant.id, "employees", current or 0)
            row = Employee(**payload.model_dump()); tdb.add(row); await tdb.commit(); await tdb.refresh(row)
            await write_audit(db, context.user.id, "admin.employee.created", "employee", str(row.id), {"tenant_id": str(tenant.id)}); await db.commit(); return _employee(row, True)
    finally: await engine.dispose()

@router.get("/{company_code}/employees/{employee_id}")
async def employee_profile(company_code: str, employee_id: uuid.UUID, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            row = await tdb.get(Employee, employee_id)
            if not row: raise HTTPException(404, "Employee not found")
            return _employee(row, context.user.role in {UserRole.COMPANY_OWNER, UserRole.HR_ADMIN})
    finally: await engine.dispose()

@router.patch("/{company_code}/employees/{employee_id}")
async def update_employee(company_code: str, employee_id: uuid.UUID, payload: EmployeeInput, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            row = await tdb.get(Employee, employee_id)
            if not row: raise HTTPException(404, "Employee not found")
            for k, v in payload.model_dump().items(): setattr(row, k, v)
            await tdb.commit(); await tdb.refresh(row); await write_audit(db, context.user.id, "admin.employee.updated", "employee", str(row.id)); await db.commit(); return _employee(row, True)
    finally: await engine.dispose()

@router.post("/{company_code}/employees/{employee_id}/archive")
async def archive_employee(company_code: str, employee_id: uuid.UUID, restore: bool = False, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            row = await tdb.get(Employee, employee_id)
            if not row: raise HTTPException(404, "Employee not found")
            row.archived_at = "" if restore else datetime.now(UTC).isoformat(); row.employment_status = "active" if restore else "archived"; await tdb.commit(); return _employee(row)
    finally: await engine.dispose()

@router.put("/{company_code}/employees/{employee_id}/photo")
async def upload_employee_photo(company_code: str, employee_id: uuid.UUID, payload: PhotoInput, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            row = await tdb.get(Employee, employee_id)
            if not row: raise HTTPException(404, "Employee not found")
            try: row.photo_data = base64.b64decode(payload.content_base64, validate=True)
            except Exception as exc: raise HTTPException(400, "Invalid base64 image") from exc
            row.photo_content_type = payload.content_type; await tdb.commit(); return {"employee_id": str(row.id), "uploaded": True}
    finally: await engine.dispose()

async def _named_list(company_code, model, context, db):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb: return [{"id": str(x.id), "name": x.name, **({"address": x.address, "manager_user_id": str(x.manager_user_id) if x.manager_user_id else None} if isinstance(x, Branch) else {})} for x in (await tdb.scalars(select(model).order_by(model.name))).all()]
    finally: await engine.dispose()

@router.get("/{company_code}/departments")
async def list_departments(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)): return await _named_list(company_code, Department, context, db)
@router.get("/{company_code}/positions")
async def list_positions(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)): return await _named_list(company_code, Position, context, db)
@router.get("/{company_code}/branches")
async def list_branches(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)): return await _named_list(company_code, Branch, context, db)

@router.post("/{company_code}/departments", status_code=201)
async def create_department(company_code: str, payload: NamedCreate, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    return await _create_named(company_code, Department, payload.model_dump(), "departments", context, db)
@router.post("/{company_code}/positions", status_code=201)
async def create_position(company_code: str, payload: PositionCreate, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    return await _create_named(company_code, Position, payload.model_dump(), "positions", context, db)
@router.post("/{company_code}/branches", status_code=201)
async def create_branch(company_code: str, payload: BranchCreate, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    return await _create_named(company_code, Branch, payload.model_dump(), "branches", context, db)

async def _create_named(company_code, model, values, limit_key, context, db):
    _admin(context); tenant, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            count = await tdb.scalar(select(func.count(model.id))); await require_limit(db, tenant.id, limit_key, count or 0)
            row = model(**values); tdb.add(row); await tdb.commit(); await tdb.refresh(row); return {"id": str(row.id), "name": row.name}
    finally: await engine.dispose()

@router.get("/{company_code}/entitlements")
async def entitlements(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); tenant, _, _ = await _ctx(company_code, context, db); e = await effective_entitlements(db, tenant.id); return {"features": e.features, "limits": e.limits, "monthly_price": e.monthly_price, "currency": e.currency}

@router.get("/{company_code}/settings")
async def settings(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb: return {x.key: x.value for x in (await tdb.scalars(select(TenantSetting))).all()}
    finally: await engine.dispose()

@router.put("/{company_code}/settings")
async def update_setting(company_code: str, payload: SettingInput, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    if context.user.role not in {UserRole.COMPANY_OWNER, UserRole.HR_ADMIN}: raise HTTPException(403, "Settings permission required")
    _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            row = await tdb.get(TenantSetting, payload.key) or TenantSetting(key=payload.key); row.value = payload.value; tdb.add(row); await tdb.commit(); return {"key": row.key, "value": row.value}
    finally: await engine.dispose()

@router.get("/{company_code}/users")
async def users(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); tenant, _, _ = await _ctx(company_code, context, db); rows = (await db.scalars(select(User).where(User.tenant_id == tenant.id))).all(); return [{"id": str(x.id), "email": x.email, "full_name": x.full_name, "role": x.role, "is_active": x.is_active} for x in rows]

@router.patch("/{company_code}/users/{user_id}/role")
async def update_user_role(company_code: str, user_id: uuid.UUID, payload: RoleUpdate, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    if context.user.role not in {UserRole.COMPANY_OWNER, UserRole.HR_ADMIN}: raise HTTPException(403, "Role management permission required")
    tenant, _, _ = await _ctx(company_code, context, db); user = await db.scalar(select(User).where(User.id == user_id, User.tenant_id == tenant.id))
    if not user: raise HTTPException(404, "User not found")
    if payload.role == UserRole.HR_ADMIN:
        count = await db.scalar(select(func.count(User.id)).where(User.tenant_id == tenant.id, User.role == UserRole.HR_ADMIN)); await require_limit(db, tenant.id, "administrators", count or 0)
    user.role = payload.role; await db.commit(); return {"id": str(user.id), "role": user.role}

@router.post("/{company_code}/employees/import")
async def import_employees(company_code: str, request: Request, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); tenant, engine, factory = await _ctx(company_code, context, db); raw = await request.body()
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True); rows = list(workbook.active.iter_rows(values_only=True))
    except ImportError: raise HTTPException(503, "Excel import dependency is not installed")
    except Exception as exc: raise HTTPException(400, f"Invalid Excel file: {exc}") from exc
    if not rows: return {"created": 0}
    headers = [str(x or "").strip() for x in rows[0]]; created = 0
    try:
        async with factory() as tdb:
            for values in rows[1:]:
                data = dict(zip(headers, values)); row = Employee(employee_id=str(data.get("employee_id", "")).strip(), first_name=str(data.get("first_name", "")).strip(), last_name=str(data.get("last_name", "")).strip(), email=str(data.get("email", "")).strip()); tdb.add(row); created += 1
            current = await tdb.scalar(select(func.count(Employee.id)).where(Employee.archived_at == "")); ent = await effective_entitlements(db, tenant.id); max_employees = ent.limits.get("employees");
            if max_employees is not None and current and current > max_employees: raise HTTPException(403, "Employee limit exceeded")
            await tdb.commit(); return {"created": created}
    finally: await engine.dispose()

@router.get("/{company_code}/employees/export.csv")
async def export_employees(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        async with factory() as tdb:
            output = io.StringIO(); writer = csv.writer(output); writer.writerow(["employee_id", "first_name", "last_name", "email", "status"])
            for x in (await tdb.scalars(select(Employee).order_by(Employee.last_name))).all(): writer.writerow([x.employee_id, x.first_name, x.last_name, x.email, x.employment_status])
            return {"filename": "employees.csv", "content_base64": base64.b64encode(output.getvalue().encode()).decode()}
    finally: await engine.dispose()

@router.get("/{company_code}/employees/export.xlsx")
async def export_employees_xlsx(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        import openpyxl
        async with factory() as tdb:
            workbook = openpyxl.Workbook(); sheet = workbook.active; sheet.append(["employee_id", "first_name", "last_name", "email", "status"])
            for x in (await tdb.scalars(select(Employee).order_by(Employee.last_name))).all(): sheet.append([x.employee_id, x.first_name, x.last_name, x.email, x.employment_status])
            output = io.BytesIO(); workbook.save(output); return {"filename": "employees.xlsx", "content_base64": base64.b64encode(output.getvalue()).decode()}
    finally: await engine.dispose()

@router.get("/{company_code}/employees/export.pdf")
async def export_employees_pdf(company_code: str, context: AuthContext = Depends(require_tenant), db: AsyncSession = Depends(get_db)):
    _admin(context); _, engine, factory = await _ctx(company_code, context, db)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen.canvas import Canvas
        async with factory() as tdb:
            output = io.BytesIO(); canvas = Canvas(output, pagesize=A4); canvas.setTitle("HR-OTO Employees"); y = 800; canvas.drawString(40, y, "HR-OTO Employees"); y -= 24
            for x in (await tdb.scalars(select(Employee).order_by(Employee.last_name))).all():
                canvas.drawString(40, y, f"{x.employee_id}  {x.first_name} {x.last_name}  {x.email}"); y -= 16
                if y < 40: canvas.showPage(); y = 800
            canvas.save(); return {"filename": "employees.pdf", "content_base64": base64.b64encode(output.getvalue()).decode()}
    finally: await engine.dispose()
